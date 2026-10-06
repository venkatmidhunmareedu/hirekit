"""Build `seed/resumes/<id>.pdf|.docx` from `seed/resumes/<id>.txt` (ADR-0012, dev only).

Deterministic: fixed metadata and dates, so a rebuild gives identical bytes. `app/` never
imports this module or its libraries. Run: `uv run python -m scripts.build_seed_resumes`.
"""

import hashlib
import io
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import structlog
from docx import Document
from fpdf import FPDF

log = structlog.get_logger()

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
LAYOUTS = ("single", "sectioned", "bullets", "table")
FIXED_DATE = datetime(2026, 1, 1, tzinfo=UTC)
SEED_DIR = Path(__file__).resolve().parents[2] / "seed" / "resumes"

type Line = tuple[str, list[str]]  # (kind, cells): kind is heading, bullet, row or text


def kind_for(resume_id: str) -> str:
    """pdf or docx by the parity of a stable hash of the id."""
    return "pdf" if hashlib.sha256(resume_id.encode()).digest()[0] % 2 == 0 else "docx"


def _parse(text: str, layout: str) -> list[Line]:
    out: list[Line] = []
    for raw in text.splitlines():
        line = raw.strip()
        if layout == "sectioned" and line and (line.isupper() or line.endswith(":")):
            out.append(("heading", [line]))
        elif layout == "bullets" and line.startswith("- "):
            out.append(("bullet", [line[2:]]))
        elif layout == "table" and " | " in line:
            out.append(("row", [c.strip() for c in line.split("|")]))
        else:
            out.append(("text", [line]))
    return out


def _pdf(lines: list[Line]) -> bytes:
    pdf = FPDF()
    pdf.set_compression(False)
    pdf.set_creation_date(FIXED_DATE)
    pdf.set_title("Resume")
    pdf.set_producer("hirekit-seed")
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for kind, cells in lines:
        if kind == "heading":
            pdf.set_font("Helvetica", "B", 13)
            pdf.multi_cell(0, 7, cells[0], new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", size=11)
        elif kind == "bullet":
            pdf.multi_cell(0, 6, "-  " + cells[0], new_x="LMARGIN", new_y="NEXT")
        elif kind == "row":
            width = pdf.epw / len(cells)
            for cell in cells:
                pdf.cell(width, 6, cell, border=1)
            pdf.ln()
        else:
            pdf.multi_cell(0, 6, cells[0], new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _docx(lines: list[Line]) -> bytes:
    doc = Document()
    props = doc.core_properties
    props.author = "hirekit-seed"
    props.title = "Resume"
    props.created = props.modified = props.last_printed = FIXED_DATE
    props.last_modified_by = "hirekit-seed"
    rows: list[list[str]] = []
    for kind, cells in [*lines, ("end", [])]:
        if kind == "row":
            rows.append(cells)
            continue
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]), style="Table Grid")
            for r, row in enumerate(rows):
                for c, cell in enumerate(row[: len(rows[0])]):
                    table.cell(r, c).text = cell
            rows = []
        if kind == "heading":
            doc.add_heading(cells[0], level=2)
        elif kind == "bullet":
            doc.add_paragraph(cells[0], style="List Bullet")
        elif kind == "text":
            doc.add_paragraph(cells[0])
    raw = io.BytesIO()
    doc.save(raw)
    # python-docx stamps zip members with the current time; rewrite them with a fixed one.
    out = io.BytesIO()
    with (
        zipfile.ZipFile(raw) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for name in src.namelist():
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, src.read(name))
    return out.getvalue()


def build(text: str, layout: str, kind: str) -> bytes:
    if layout not in LAYOUTS:
        raise ValueError(f"unknown layout {layout!r}")
    if kind not in ("pdf", "docx"):
        raise ValueError(f"unknown kind {kind!r}")
    lines = _parse(text, layout)
    return _pdf(lines) if kind == "pdf" else _docx(lines)


def build_all(directory: Path) -> list[Path]:
    """Build every `<id>.txt` in `directory`; layout cycles by a stable hash of the id."""
    written: list[Path] = []
    for source in sorted(directory.glob("*.txt")):
        resume_id = source.stem
        kind = kind_for(resume_id)
        layout = LAYOUTS[hashlib.sha256(resume_id.encode()).digest()[1] % len(LAYOUTS)]
        target = source.with_suffix(f".{kind}")
        target.write_bytes(build(source.read_text(encoding="utf-8"), layout, kind))
        log.info("seed_resume_built", resume_id=resume_id, layout=layout, kind=kind)
        written.append(target)
    return written


def main() -> int:
    written = build_all(SEED_DIR)
    log.info("seed_resumes_done", count=len(written))
    return 0 if written else 1


if __name__ == "__main__":
    sys.exit(main())
